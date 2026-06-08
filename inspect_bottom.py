import openpyxl

wb = openpyxl.load_workbook('financial-anaylisis.xlsx')
ws = wb['sea \u8d22\u62a5V2']

non_empty_rows = []
for r in range(53, ws.max_row + 1):
    row_vals = [ws.cell(row=r, column=col).value for col in range(1, 10)]
    if any(val is not None for val in row_vals):
        non_empty_rows.append((r, row_vals))

print(f"Total non-empty rows after row 52: {len(non_empty_rows)}")
for r, vals in non_empty_rows[:30]:
    print(f"Row {r:03d}: {vals}")
