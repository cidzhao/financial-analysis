import openpyxl
import json

wb = openpyxl.load_workbook('c:/Users/zhaoxiaoyu/Documents/finical/finical-report/sea/sea-anaylisis.xlsx', data_only=True)
ws = wb.worksheets[0]

data = {}
# Read header (Quarters)
headers = []
for col in range(2, ws.max_column + 1):
    val = ws.cell(row=1, column=col).value
    if val:
        headers.append(str(val))

data['quarters'] = headers

# Read row data
for row in range(2, ws.max_row + 1):
    row_name = ws.cell(row=row, column=1).value
    if row_name:
        row_name = str(row_name).strip()
        row_data = []
        for col in range(2, len(headers) + 2):
            val = ws.cell(row=row, column=col).value
            row_data.append(val)
        data[row_name] = row_data

with open('c:/Users/zhaoxiaoyu/Documents/finical/finical-report/sea/extracted_data.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, indent=2)

print("Data extracted.")
